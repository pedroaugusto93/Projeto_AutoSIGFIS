# logger.py
# -*- coding: utf-8 -*-
"""
Camada de observabilidade do AutoSIGFIS.

Fornece:
  - setup_logging()      -> configura log em console + arquivo (logs/autosigfis_*.log)
  - get_logger(nome)     -> logger filho para cada módulo
  - dump_failure(driver) -> salva screenshot + page_source no momento da falha
  - ExecutionReport      -> acumula o status por registro e gera o extrato .xlsx

Mensagens usam apenas ASCII para não quebrar o console do Windows (cp1252).
O arquivo de log é UTF-8 e guarda o detalhe campo a campo (nível DEBUG).
"""

import sys
import logging
from datetime import datetime
from pathlib import Path

LOG_DIR = Path(__file__).resolve().parent / "logs"
LOG_DIR.mkdir(exist_ok=True)

LOG_FILE = LOG_DIR / f"autosigfis_{datetime.now():%Y%m%d_%H%M%S}.log"

_ROOT_NAME = "autosigfis"


def setup_logging(console_level: int = logging.INFO) -> logging.Logger:
    """Configura (uma única vez) o logger raiz do projeto."""
    logger = logging.getLogger(_ROOT_NAME)
    if logger.handlers:                      # já configurado
        return logger

    logger.setLevel(logging.DEBUG)
    fmt = logging.Formatter(
        "%(asctime)s | %(levelname)-7s | %(name)s | %(message)s",
        datefmt="%Y-%m-%d %H:%M:%S",
    )

    # Arquivo: tudo (DEBUG), inclusive cada campo preenchido
    fh = logging.FileHandler(LOG_FILE, encoding="utf-8")
    fh.setLevel(logging.DEBUG)
    fh.setFormatter(fmt)
    logger.addHandler(fh)

    # Console: resumo (INFO por padrão)
    try:
        sys.stdout.reconfigure(encoding="utf-8")   # Python 3.7+; evita UnicodeEncodeError
    except Exception:
        pass
    ch = logging.StreamHandler(sys.stdout)
    ch.setLevel(console_level)
    ch.setFormatter(fmt)
    logger.addHandler(ch)

    logger.propagate = False
    return logger


def get_logger(nome: str = _ROOT_NAME) -> logging.Logger:
    """Retorna o logger raiz ou um filho (ex.: get_logger('dados_basicos'))."""
    root = logging.getLogger(_ROOT_NAME)
    return root if nome in (_ROOT_NAME, "", None) else root.getChild(nome)


def dump_failure(driver, prefixo: str = "falha") -> None:
    """Salva screenshot + HTML da página no momento de uma falha (para diagnóstico)."""
    log = get_logger()
    ts = datetime.now().strftime("%H%M%S")
    try:
        png = LOG_DIR / f"{prefixo}_{ts}.png"
        driver.save_screenshot(str(png))
        log.info("Screenshot da falha: %s", png.name)
    except Exception as e:
        log.warning("Nao foi possivel salvar screenshot: %s", e)
    try:
        html = LOG_DIR / f"{prefixo}_{ts}.html"
        with open(html, "w", encoding="utf-8") as f:
            f.write(driver.page_source)
        log.info("HTML da falha: %s", html.name)
    except Exception as e:
        log.warning("Nao foi possivel salvar page_source: %s", e)


class ExecutionReport:
    """Acumula um resumo por registro e gera o extrato .xlsx ao final da execução."""

    COLUNAS = [
        "registro", "processo", "nome_fornecedor", "cnpj_fornecedor",
        "valor", "num_empenho", "dispensa_id", "status", "perc_conclusao",
        "validacoes", "ultima_aba", "etapa_falha", "erro_tipo", "erro_msg",
        "duracao_s", "inicio",
    ]

    def __init__(self):
        self.rows = []

    def add(self, **kwargs):
        self.rows.append(kwargs)

    def resumo(self):
        ok = sum(1 for r in self.rows if r.get("status") == "OK")
        return ok, len(self.rows) - ok, len(self.rows)

    def save(self) -> Path:
        """Grava o extrato em logs/extrato_*.xlsx. Retorna o caminho."""
        log = get_logger()
        path = LOG_DIR / f"extrato_{datetime.now():%Y%m%d_%H%M%S}.xlsx"
        try:
            import pandas as pd
            df = pd.DataFrame(self.rows)
            # Garante ordem de colunas, mesmo que alguma falte
            cols = [c for c in self.COLUNAS if c in df.columns]
            extras = [c for c in df.columns if c not in cols]
            df = df[cols + extras]
            df.to_excel(path, index=False)
            self._polir(path)
        except Exception as e:
            log.error("Falha ao gerar extrato .xlsx: %s", e)
        return path

    @staticmethod
    def _polir(path: Path) -> None:
        """Ajuste cosmético leve do extrato (largura de coluna, cabeçalho, freeze)."""
        try:
            from openpyxl import load_workbook
            from openpyxl.styles import Font, PatternFill, Alignment
            wb = load_workbook(path)
            ws = wb.active
            header_fill = PatternFill("solid", fgColor="1F4E78")
            header_font = Font(color="FFFFFF", bold=True)
            for cell in ws[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(vertical="center")
            for col in ws.columns:
                largura = max((len(str(c.value)) for c in col if c.value is not None), default=10)
                ws.column_dimensions[col[0].column_letter].width = min(largura + 3, 60)
            ws.freeze_panes = "A2"
            wb.save(path)
        except Exception:
            pass  # cosmético: nunca derruba a execução
